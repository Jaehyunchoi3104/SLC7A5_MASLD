# Repository paths; inputs are read-only, new files stay in results/.
from pathlib import Path
import os
import sys
_start = Path(os.environ.get("SLC7A5_REPO_ROOT", str(Path(__file__).resolve().parent))).resolve()
_repo = next((p for p in [_start, *_start.parents] if (p / "slc7a5_paths.py").is_file()), None)
if _repo is None:
    raise RuntimeError("Run inside the repository or set SLC7A5_REPO_ROOT.")
if str(_repo) not in sys.path:
    sys.path.insert(0, str(_repo))
from slc7a5_paths import dataset_path, output_path, setup_notebook

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.styles.numbers import FORMAT_NUMBER_00

# ── Cluster annotation ──────────────────────────────────────────────────────
ANNOTATIONS = {
    "mono0": {
        "label": "Classical Monocyte",
        "key_markers": "FCN1, VCAN, CD14, S100A8, S100A9, S100A12, LYZ, MNDA",
        "function": "Inflammatory response, innate immune response, phagocytosis",
    },
    "mono1": {
        "label": "Kupffer Cell",
        "key_markers": "CD163, MARCO, CD5L, C1QA, C1QB, C1QC, VCAM1, CXCL12",
        "function": "Lysosome-mediated degradation, endocytosis, tissue homeostasis",
    },
    "mono2": {
        "label": "cDC2",
        "key_markers": "CD1C, CLEC10A, FCER1A, HLA-DQA1, HLA-DQB1, HLA-DPA1, HLA-DPB1, CD74",
        "function": "MHC class II antigen presentation, hematopoietic cell differentiation",
    },
    "mono3": {
        "label": "Non-classical Monocyte",
        "key_markers": "FCGR3A (CD16), LST1, IFITM2, IFITM3, RHOC, CORO1A, FGR, SPN",
        "function": "Vascular patrolling, cytoskeletal regulation, NF-κB signaling",
    },
    "mono4": {
        "label": "LAM (Lipid-Associated Macrophage)",
        "key_markers": "TREM2, GPNMB, LGALS3, APOE, FABP5, CD9, CD68, APOC1",
        "function": "Lysosomal lipid processing, phagosome activity, MHC II presentation",
    },
    "mono5": {
        "label": "Special Monocyte",
        "key_markers": "ANPEP (CD13), ITGAX (CD11c), CD300E, FCN1, VCAN, SLC7A5",
        "function": "Transcriptional activation, NF-κB signaling, intracellular signal transduction (CD14-maintained, CD16-low)",
    },
    "mono7": {
        "label": "Monocyte-derived Kupffer Cell",
        "key_markers": "C1QA, C1QB, C1QC, AIF1 (IBA1), FTL, FTH1, MS4A6A, FCER1G",
        "function": "Complement activation, innate immune response, iron metabolism",
    },
    "mono8": {
        "label": "cDC1",
        "key_markers": "CLEC9A (DNGR-1), IRF8, IDO1, DNASE1L3, HLA-DQA1, HLA-DPB1",
        "function": "Antigen processing and presentation (MHC I & II), cross-presentation",
    },
    "mono9": {
        "label": "Migratory DC",
        "key_markers": "LAMP3 (CD208), CD83, CD40, IL4I1, CXCL10, GBP1, ISG20",
        "function": "NF-κB signaling, interferon response, DC maturation and migration",
    },
}

CLUSTER_ORDER = ["mono0","mono1","mono2","mono3","mono4","mono5","mono7","mono8","mono9"]

# ── Color palette ────────────────────────────────────────────────────────────
HEADER_FILLS = {
    "mono0": "2E75B6",  # blue - classical monocyte
    "mono1": "C55A11",  # brown - KC
    "mono2": "70AD47",  # green - cDC2
    "mono3": "4472C4",  # cornflower - non-classical
    "mono4": "7030A0",  # purple - LAM
    "mono5": "00B0F0",  # cyan - special monocyte
    "mono7": "FF7C00",  # orange - moDK
    "mono8": "375623",  # dark green - cDC1
    "mono9": "843C0C",  # dark red - migratory DC
}
ROW_ALT_FILL = "F2F2F2"

# ── Helper styles ────────────────────────────────────────────────────────────
def header_font(color="FFFFFF"):
    return Font(name="Calibri", bold=True, color=color, size=11)

def body_font():
    return Font(name="Calibri", size=10)

def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=False)

def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=False)

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def set_col_width(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width

# ── Overview sheet ───────────────────────────────────────────────────────────
def make_overview(wb, df):
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells("A1:F1")
    tc = ws["A1"]
    tc.value = "Supplementary Table – Monocyte/Macrophage Cluster DEG Analysis"
    tc.font = Font(name="Calibri", bold=True, size=13, color="1F3864")
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Sub-header
    ws.merge_cells("A2:F2")
    sc = ws["A2"]
    sc.value = ("scRNA-seq differential expression results for each mononuclear phagocyte cluster. "
                "Significant DEGs (adjusted p-value < 0.05, |log₂FC| > 0) are shown per cluster sheet.")
    sc.font = Font(name="Calibri", italic=True, size=10, color="595959")
    sc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30

    ws.row_dimensions[3].height = 6  # spacer

    # Column headers row 4
    headers = ["Cluster ID", "Cell Type Annotation", "# Sig. DEGs (padj<0.05)",
               "Key Marker Genes", "Biological Function", "Sheet"]
    col_widths = [14, 32, 22, 50, 60, 22]

    for i, (h, w) in enumerate(zip(headers, col_widths), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = center()
        c.border = thin_border()
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[4].height = 18

    # Data rows
    for row_i, cluster in enumerate(CLUSTER_ORDER, start=5):
        ann = ANNOTATIONS[cluster]
        sig_count = len(df[(df["group"] == cluster) & (df["pvals_adj"] < 0.05)])
        sheet_name = f"{cluster.replace('mono','M')} – {ann['label']}"

        values = [
            cluster.replace("mono", "M"),
            ann["label"],
            sig_count,
            ann["key_markers"],
            ann["function"],
            sheet_name,
        ]
        fill_color = ROW_ALT_FILL if row_i % 2 == 0 else "FFFFFF"
        for col_i, val in enumerate(values, start=1):
            c = ws.cell(row=row_i, column=col_i, value=val)
            c.font = body_font()
            c.fill = PatternFill("solid", fgColor=fill_color)
            c.border = thin_border()
            c.alignment = left() if col_i > 1 else center()
        ws.row_dimensions[row_i].height = 16

    # Color the cluster ID cell with its own color
    for row_i, cluster in enumerate(CLUSTER_ORDER, start=5):
        c = ws.cell(row=row_i, column=1)
        c.fill = PatternFill("solid", fgColor=HEADER_FILLS[cluster])
        c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        c.alignment = center()

    ws.freeze_panes = "A5"

# ── Per-cluster DEG sheet ────────────────────────────────────────────────────
def make_cluster_sheet(wb, df, cluster):
    ann = ANNOTATIONS[cluster]
    sheet_name = f"{cluster.replace('mono','M')} – {ann['label']}"
    # Excel sheet name max 31 chars
    if len(sheet_name) > 31:
        sheet_name = sheet_name[:31]

    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_view.showGridLines = False

    hdr_color = HEADER_FILLS[cluster]

    # Row 1: Cluster annotation header
    ws.merge_cells("A1:I1")
    c1 = ws["A1"]
    c1.value = f"{cluster.replace('mono','M')}  |  {ann['label']}"
    c1.font = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    c1.fill = PatternFill("solid", fgColor=hdr_color)
    c1.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    # Row 2: Key markers
    ws.merge_cells("A2:I2")
    c2 = ws["A2"]
    c2.value = f"Key markers: {ann['key_markers']}"
    c2.font = Font(name="Calibri", italic=True, size=10, color="404040")
    c2.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 18

    # Row 3: Function
    ws.merge_cells("A3:I3")
    c3 = ws["A3"]
    c3.value = f"Function: {ann['function']}"
    c3.font = Font(name="Calibri", italic=True, size=10, color="404040")
    c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 18

    ws.row_dimensions[4].height = 6  # spacer

    # Row 5: Column headers
    col_defs = [
        ("Rank",            8),
        ("Gene",           14),
        ("Score",          12),
        ("Log₂FC",         11),
        ("p-value",        16),
        ("Adj. p-value",   16),
        ("% Expr. (cluster)", 18),
        ("% Expr. (others)",  18),
        ("Significance",   14),
    ]
    for col_i, (name, width) in enumerate(col_defs, start=1):
        c = ws.cell(row=5, column=col_i, value=name)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=hdr_color)
        c.alignment = center()
        c.border = thin_border()
        ws.column_dimensions[get_column_letter(col_i)].width = width
    ws.row_dimensions[5].height = 18

    # Filter: padj < 0.05, sort by score desc
    sub = df[df["group"] == cluster].copy()
    sig = sub[sub["pvals_adj"] < 0.05].sort_values("scores", ascending=False).reset_index(drop=True)

    for row_i, (_, row) in enumerate(sig.iterrows(), start=1):
        excel_row = row_i + 5

        # Significance label
        padj = row["pvals_adj"]
        if padj == 0:
            sig_label = "****"
        elif padj < 0.001:
            sig_label = "***"
        elif padj < 0.01:
            sig_label = "**"
        elif padj < 0.05:
            sig_label = "*"
        else:
            sig_label = "ns"

        values = [
            row_i,
            row["names"],
            round(float(row["scores"]), 4),
            round(float(row["logfoldchanges"]), 4),
            float(row["pvals"]),
            float(row["pvals_adj"]),
            round(float(row["pct_nz_group"]) * 100, 2),
            round(float(row["pct_nz_reference"]) * 100, 2),
            sig_label,
        ]

        fill_color = ROW_ALT_FILL if row_i % 2 == 0 else "FFFFFF"
        for col_i, val in enumerate(values, start=1):
            c = ws.cell(row=excel_row, column=col_i, value=val)
            c.font = body_font()
            c.fill = PatternFill("solid", fgColor=fill_color)
            c.border = thin_border()
            # Alignment
            if col_i == 2:  # Gene name
                c.alignment = left()
            elif col_i == 9:  # Significance stars
                c.alignment = center()
                c.font = Font(name="Calibri", bold=True, size=10,
                              color="C00000" if val not in ("ns", "") else "595959")
            else:
                c.alignment = center()

            # Scientific notation for p-values
            if col_i in (5, 6):
                c.number_format = "0.00E+00"

        ws.row_dimensions[excel_row].height = 14

    # Freeze panes below header
    ws.freeze_panes = "A6"

    # Total sig DEG count at bottom
    bottom_row = len(sig) + 7
    ws.merge_cells(f"A{bottom_row}:I{bottom_row}")
    bc = ws.cell(row=bottom_row, column=1,
                 value=f"Total significant DEGs (adj. p-value < 0.05): {len(sig)}")
    bc.font = Font(name="Calibri", bold=True, italic=True, size=10, color=hdr_color)
    bc.alignment = left()

# ── Main ─────────────────────────────────────────────────────────────────────
def main():
    setup_notebook("monomac_table")
    csv_path = dataset_path("monomac_marker_deg")
    out_path = output_path("02_SLC7A5_mono_Journal/05_Figure/monomacDEG/Suppl_table_monoDEG.xlsx")

    print("Reading CSV...")
    df = pd.read_csv(csv_path, index_col=0)

    print("Building workbook...")
    wb = Workbook()

    make_overview(wb, df)
    for cluster in CLUSTER_ORDER:
        print(f"  Sheet: {cluster}")
        make_cluster_sheet(wb, df, cluster)

    print(f"Saving to {out_path}")
    wb.save(out_path)
    print("Done.")

if __name__ == "__main__":
    main()
